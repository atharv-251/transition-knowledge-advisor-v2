from __future__ import annotations

import csv
import hashlib
import logging
from abc import ABC, abstractmethod
from datetime import date, datetime, time
from pathlib import Path
from typing import Iterable
from zoneinfo import ZoneInfo

from pydantic import ValidationError

from app.kt_scheduler.models import KtScheduleItem

logger = logging.getLogger("kt_scheduler.input_sources")


class PlanInputSource(ABC):
    """Interface for any current/future KT plan source."""

    @abstractmethod
    def load(self) -> list[KtScheduleItem]:
        """Return validated KT schedule items."""


class CsvPlanInputSource(PlanInputSource):
    """CSV input source for the current exported KT plan file."""

    def __init__(self, path: str | Path, timezone_name: str) -> None:
        self.path = Path(path)
        self.timezone = ZoneInfo(timezone_name)

    def load(self) -> list[KtScheduleItem]:
        if not self.path.exists():
            raise FileNotFoundError(f"KT schedule input was not found: {self.path}")

        items: list[KtScheduleItem] = []
        with self.path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            for row_number, row in enumerate(reader, start=2):
                if _is_blank_row(row.values()):
                    continue
                try:
                    items.append(_row_to_item(row, row_number, self.timezone))
                except (ValueError, ValidationError) as exc:
                    logger.warning(
                        "Skipping invalid KT schedule row %s from %s: %s",
                        row_number,
                        self.path,
                        exc,
                    )
        return items


class ExcelPlanInputSource(PlanInputSource):
    """XLSX input source kept for the real Excel version of the same plan."""

    def __init__(self, path: str | Path, timezone_name: str) -> None:
        self.path = Path(path)
        self.timezone = ZoneInfo(timezone_name)

    def load(self) -> list[KtScheduleItem]:
        if not self.path.exists():
            raise FileNotFoundError(f"KT schedule input was not found: {self.path}")

        from openpyxl import load_workbook

        workbook = load_workbook(self.path, data_only=True, read_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        items: list[KtScheduleItem] = []
        for row_number, values in enumerate(rows, start=2):
            row = {
                header: "" if value is None else str(value).strip()
                for header, value in zip(headers, values)
            }
            if _is_blank_row(row.values()):
                continue
            try:
                items.append(_row_to_item(row, row_number, self.timezone))
            except (ValueError, ValidationError) as exc:
                logger.warning(
                    "Skipping invalid KT schedule row %s from %s: %s",
                    row_number,
                    self.path,
                    exc,
                )
        return items


class BotApiPlanInputSource(PlanInputSource):
    """Future adapter placeholder for another bot/API providing KT schedules."""

    def load(self) -> list[KtScheduleItem]:
        raise NotImplementedError(
            "BotApiPlanInputSource is intentionally a placeholder. Implement "
            "this adapter when the KT plan source bot/API contract is ready."
        )


def input_source_for(path: str | Path, timezone_name: str) -> PlanInputSource:
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        return CsvPlanInputSource(path, timezone_name)
    if suffix in (".xlsx", ".xlsm"):
        return ExcelPlanInputSource(path, timezone_name)
    raise ValueError(
        "Unsupported KT schedule input file type. Expected .csv, .xlsx or .xlsm."
    )


def _row_to_item(row: dict[str, str], row_number: int, timezone: ZoneInfo) -> KtScheduleItem:
    title = _required(row, "Session Title", row_number)
    scheduled_date = _parse_date(_required(row, "Scheduled Date", row_number))
    start_time = _parse_time(_required(row, "Start Time", row_number))
    end_time = _parse_time(_required(row, "End Time", row_number))

    start_at = datetime.combine(scheduled_date, start_time, tzinfo=timezone)
    end_at = datetime.combine(scheduled_date, end_time, tzinfo=timezone)
    if end_at <= start_at:
        raise ValueError(f"Row {row_number}: End Time must be after Start Time")

    required_attendees = _split_emails(row.get("Required Attendees", ""))
    optional_attendees = _split_emails(row.get("Optional Attendees", ""))

    fingerprint = hashlib.sha256(
        "|".join([
            title,
            start_at.isoformat(),
            end_at.isoformat(),
            ";".join(required_attendees),
        ]).encode("utf-8")
    ).hexdigest()[:12]

    return KtScheduleItem(
        source_id=f"csv-row-{row_number}-{fingerprint}",
        session_title=title,
        level=(row.get("Level") or "").strip(),
        start_at=start_at,
        end_at=end_at,
        duration_hours=_optional_float(row.get("Duration Hours")),
        delivery_mode=(row.get("Delivery Mode") or "").strip(),
        assigned_sme=(row.get("Assigned SME") or "").strip(),
        assigned_receiver=(row.get("Assigned Receiver") or "").strip(),
        required_attendees=required_attendees,
        optional_attendees=optional_attendees,
        status=(row.get("Status") or "").strip(),
        conflicts=(row.get("Conflicts") or "").strip(),
        description=_build_description(row),
        agenda_topics=[title],
    )


def _build_description(row: dict[str, str]) -> str:
    return "\n".join([
        f"KT Level: {(row.get('Level') or '').strip()}",
        f"Delivery Mode: {(row.get('Delivery Mode') or '').strip()}",
        f"Assigned SME: {(row.get('Assigned SME') or '').strip()}",
        f"Assigned Receiver: {(row.get('Assigned Receiver') or '').strip()}",
        "",
        "Agenda:",
        f"- {(row.get('Session Title') or '').strip()}",
    ]).strip()


def _required(row: dict[str, str], column: str, row_number: int) -> str:
    value = (row.get(column) or "").strip()
    if not value:
        raise ValueError(f"Row {row_number}: missing required column {column!r}")
    return value


def _parse_date(value: str) -> date:
    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(value.strip(), fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Invalid date value: {value!r}")


def _parse_time(value: str) -> time:
    value = value.strip()
    for fmt in ("%H:%M", "%I:%M %p", "%I:%M%p"):
        try:
            return datetime.strptime(value, fmt).time()
        except ValueError:
            continue
    raise ValueError(f"Invalid time value: {value!r}")


def _split_emails(value: str) -> list[str]:
    raw_items = value.replace(",", ";").split(";")
    return [item.strip() for item in raw_items if item.strip()]


def _optional_float(value: str | None) -> float | None:
    if value is None or not str(value).strip():
        return None
    return float(str(value).strip())


def _is_blank_row(values: Iterable[object]) -> bool:
    return all(value is None or not str(value).strip() for value in values)

