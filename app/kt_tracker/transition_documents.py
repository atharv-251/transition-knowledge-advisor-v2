"""Read KT Planner transition documents from the local Transition_Docs folder."""

from __future__ import annotations

import csv
import json
import os
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

DEFAULT_TRANSITIONS_ROOT = Path(__file__).resolve().parents[3] / "Transition_Docs"
LATEST_TRANSITION_FILE = ".latest_transition.json"


def transitions_root() -> Path:
    configured = os.getenv("KT_TRANSITIONS_ROOT", "").strip()
    return Path(configured) if configured else DEFAULT_TRANSITIONS_ROOT


def _transition_id(transition_name: str) -> str:
    value = re.sub(r"[^A-Za-z0-9_-]+", "_", transition_name.strip()).strip("_-")
    if not value:
        raise ValueError("Transition name must contain letters or numbers.")
    return value


def _safe_filename(filename: str, expected_suffix: str) -> str:
    path = Path(filename or "")
    if path.suffix.lower() != expected_suffix:
        raise ValueError(f"Expected a {expected_suffix} file.")
    name = re.sub(r"[^A-Za-z0-9._-]+", "_", path.name)
    if not name:
        raise ValueError("Uploaded file name is invalid.")
    return name


def save_uploaded_transition(
    *,
    transition_name: str,
    master_plan_name: str,
    master_plan_content: bytes,
    schedule_name: str,
    schedule_content: bytes,
) -> dict[str, str]:
    """Save an uploaded Planner package and mark it as the latest transition."""
    transition_id = _transition_id(transition_name)
    master_plan_filename = _safe_filename(master_plan_name, ".xlsx")
    schedule_filename = _safe_filename(schedule_name, ".csv")
    root = transitions_root()
    transition_dir = root / transition_id
    master_plan_dir = transition_dir / "Master_Plan"
    schedule_dir = transition_dir / "Schedule"
    master_plan_dir.mkdir(parents=True, exist_ok=True)
    schedule_dir.mkdir(parents=True, exist_ok=True)

    for existing in _files(master_plan_dir, "*.xlsx"):
        existing.unlink()
    for existing in _files(schedule_dir, "*.csv"):
        existing.unlink()
    (master_plan_dir / master_plan_filename).write_bytes(master_plan_content)
    (schedule_dir / schedule_filename).write_bytes(schedule_content)

    (root / LATEST_TRANSITION_FILE).write_text(
        json.dumps({"transition_id": transition_id}), encoding="utf-8"
    )
    return {
        "id": transition_id,
        "name": transition_name.strip(),
        "master_plan": master_plan_filename,
        "schedule": schedule_filename,
        "teams_transcript": "",
    }


def save_teams_transcript(
    *,
    transition_name: str,
    transcript_name: str,
    transcript_content: bytes,
) -> str:
    """Attach one Teams transcript to an existing local transition."""
    _document_paths(transition_name)
    transcript_filename = _safe_filename(transcript_name, ".vtt")
    transcript_dir = transitions_root() / transition_name / "Teams_Transcripts"
    transcript_dir.mkdir(exist_ok=True)
    for existing in _files(transcript_dir, "*.vtt"):
        existing.unlink()
    (transcript_dir / transcript_filename).write_bytes(transcript_content)
    return transcript_filename


def latest_transition_id() -> str | None:
    """Return the last uploaded transition, falling back to an existing folder."""
    root = transitions_root()
    manifest_path = root / LATEST_TRANSITION_FILE
    if manifest_path.is_file():
        try:
            transition_id = str(json.loads(manifest_path.read_text(encoding="utf-8"))["transition_id"])
            _document_paths(transition_id)
            return transition_id
        except (KeyError, ValueError, json.JSONDecodeError):
            pass
    transitions = list_transitions()
    return transitions[-1]["id"] if transitions else None


def _files(directory: Path, pattern: str) -> list[Path]:
    return sorted(path for path in directory.glob(pattern) if not path.name.startswith("~$"))


def _document_paths(transition_name: str) -> tuple[Path, Path]:
    root = transitions_root()
    transition_dir = root / transition_name
    if not transition_dir.is_dir() or transition_dir.parent != root:
        raise KeyError(f"Transition {transition_name!r} was not found.")

    master_plans = _files(transition_dir / "Master_Plan", "*.xlsx")
    schedules = _files(transition_dir / "Schedule", "*.csv")
    if not master_plans or not schedules:
        raise ValueError(
            f"Transition {transition_name!r} must contain one .xlsx file in "
            "Master_Plan and one .csv file in Schedule."
        )
    return master_plans[0], schedules[0]


def list_transitions() -> list[dict[str, str]]:
    root = transitions_root()
    if not root.is_dir():
        return []

    transitions = []
    for directory in sorted(path for path in root.iterdir() if path.is_dir()):
        master_plans = _files(directory / "Master_Plan", "*.xlsx")
        schedules = _files(directory / "Schedule", "*.csv")
        if master_plans or schedules:
            transcripts = _files(directory / "Teams_Transcripts", "*.vtt")
            transitions.append(
                {
                    "id": directory.name,
                    "name": directory.name.replace("_", " "),
                    "master_plan_file": master_plans[0].name if master_plans else "",
                    "schedule_file": schedules[0].name if schedules else "",
                    "teams_transcript": transcripts[0].name if transcripts else "",
                }
            )
    return transitions


def _worksheet_rows(workbook: Any, worksheet_name: str) -> list[dict[str, Any]]:
    worksheet = workbook[worksheet_name]
    rows = list(worksheet.iter_rows(values_only=True))
    header_index = next(
        (index for index, row in enumerate(rows) if any(value is not None for value in row) and index > 1),
        None,
    )
    if header_index is None:
        return []
    headers = [str(value).strip() if value is not None else "" for value in rows[header_index]]
    records = []
    for row in rows[header_index + 1 :]:
        if not any(value is not None and str(value).strip() for value in row):
            continue
        records.append(
            {
                header: _json_value(value)
                for header, value in zip(headers, row)
                if header
            }
        )
    return records


def _json_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    return value


def _executive_summary(workbook: Any) -> dict[str, str]:
    rows = _worksheet_rows(workbook, "Executive Summary")
    return {
        str(row.get("Parameter", "")): str(row.get("Transition Value", ""))
        for row in rows
        if row.get("Parameter")
    }


def _read_schedule(schedule_path: Path) -> list[dict[str, str]]:
    with schedule_path.open(encoding="utf-8-sig", newline="") as handle:
        return [dict(row) for row in csv.DictReader(handle)]


def _calendar_days(sessions: list[dict[str, str]]) -> list[dict[str, Any]]:
    grouped: dict[str, list[dict[str, str]]] = defaultdict(list)
    for session in sessions:
        date_value = session.get("Scheduled Date", "")
        grouped[date_value].append(session)

    def date_sort_key(value: str) -> datetime:
        try:
            return datetime.strptime(value, "%m/%d/%Y")
        except ValueError:
            return datetime.max

    return [
        {
            "date": date_value,
            "day": day_sessions[0].get("Day", ""),
            "sessions": day_sessions,
        }
        for date_value, day_sessions in sorted(grouped.items(), key=lambda item: date_sort_key(item[0]))
    ]


def get_transition(transition_name: str) -> dict[str, Any]:
    master_plan_path, schedule_path = _document_paths(transition_name)
    workbook = load_workbook(master_plan_path, read_only=True, data_only=True)
    try:
        master_plan = _worksheet_rows(workbook, "KT Master Plan")
        validation = _worksheet_rows(workbook, "Validation Report")
        availability = _worksheet_rows(workbook, "Availability Report")
        summary = _executive_summary(workbook)
    finally:
        workbook.close()

    sessions = _read_schedule(schedule_path)
    transcript_files = _files(
        transitions_root() / transition_name / "Teams_Transcripts", "*.vtt"
    )
    return {
        "id": transition_name,
        "name": summary.get("Transition Name", transition_name).strip(),
        "files": {
            "master_plan": master_plan_path.name,
            "schedule": schedule_path.name,
            "teams_transcript": transcript_files[0].name if transcript_files else "",
        },
        "summary": summary,
        "master_plan": master_plan,
        "schedule": sessions,
        "calendar": _calendar_days(sessions),
        "validation": validation,
        "availability": availability,
    }
